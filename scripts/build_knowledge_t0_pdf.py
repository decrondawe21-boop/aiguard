from __future__ import annotations

import os
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "dist" / "knowledge"


def default_render_script() -> Path:
    candidates = sorted(
        (Path.home() / ".codex" / "plugins" / "cache" / "openai-primary-runtime" / "documents").glob(
            "*/skills/documents/render_docx.py"
        )
    )
    return candidates[-1] if candidates else Path("render_docx.py")


RENDER_SCRIPT = Path(os.environ.get("AEGIS_RENDER_DOCX", str(default_render_script())))
PYTHON = Path(os.environ.get("AEGIS_PYTHON", sys.executable))

TITLE = RGBColor(15, 35, 63)
ACCENT = RGBColor(10, 129, 151)
INK = RGBColor(34, 39, 50)
MUTED = RGBColor(88, 99, 118)
LIGHT_FILL = "EAF6F8"
TABLE_HEADER = "DDEFF3"
CALLOUT = "F2F8FA"


def set_font(run, *, size: int, color: RGBColor = INK, bold: bool = False) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_padding(cell, value: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "bottom", "start", "end"):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_width(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def p(doc: Document, text: str = "", *, size: int = 10, color: RGBColor = INK, bold: bool = False,
      align=WD_ALIGN_PARAGRAPH.LEFT, before: int = 0, after: int = 6, line: float = 1.18):
    para = doc.add_paragraph()
    para.alignment = align
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = line
    run = para.add_run(text)
    set_font(run, size=size, color=color, bold=bold)
    return para


def heading(doc: Document, text: str, level: int = 1) -> None:
    sizes = {1: 16, 2: 13, 3: 11}
    colors = {1: ACCENT, 2: TITLE, 3: TITLE}
    before = {1: 14, 2: 10, 3: 8}
    after = {1: 7, 2: 5, 3: 4}
    para = p(
        doc,
        text,
        size=sizes[level],
        color=colors[level],
        bold=True,
        before=before[level],
        after=after[level],
    )
    para.style = f"Heading {level}"


def bullet(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.15
    run = para.add_run(text)
    set_font(run, size=10, color=INK)


def numbered(doc: Document, text: str) -> None:
    para = doc.add_paragraph(style="List Number")
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.15
    run = para.add_run(text)
    set_font(run, size=10, color=INK)


def callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(16.4)
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT)
    set_cell_padding(cell, 150)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(f"{label}: ")
    set_font(label_run, size=10, color=ACCENT, bold=True)
    body_run = paragraph.add_run(text)
    set_font(body_run, size=10, color=INK)
    doc.add_paragraph()


def table_key_values(doc: Document, rows: list[tuple[str, str]], *, widths: tuple[float, float] = (4.2, 12.2)) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [widths[0], widths[1]])
    header = table.rows[0].cells
    header[0].text = "Oblast"
    header[1].text = "Znalost T=0"
    for cell in header:
        shade_cell(cell, TABLE_HEADER)
        set_cell_padding(cell, 120)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, size=9, color=TITLE, bold=True)

    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for idx, cell in enumerate(cells):
            set_cell_padding(cell, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.12
                for run in paragraph.runs:
                    set_font(run, size=9, color=TITLE if idx == 0 else INK, bold=idx == 0)
    doc.add_paragraph()


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21.59)
    section.page_height = Cm(27.94)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)

    header = section.header.paragraphs[0]
    header.text = "AEGIS | Knowledge T=0"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.text = "aegis.d-international.eu | David Kozák | stav k 10. 5. 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in footer.runs:
        set_font(run, size=8, color=MUTED)

    return doc


def build_docx() -> Path:
    doc = setup_document()

    p(doc, "AEGIS", size=28, color=TITLE, bold=True, after=2)
    p(doc, "Knowledge T=0", size=18, color=ACCENT, bold=True, after=4)
    p(
        doc,
        "Výchozí znalostní báze projektu podle funkcí, které jsou v repozitáři aktuálně implementované.",
        size=11,
        color=MUTED,
        after=16,
    )
    table_key_values(
        doc,
        [
            ("Datum snapshotu", "10. 5. 2026"),
            ("Produkční doména", "https://aegis.d-international.eu"),
            ("Repo root", str(ROOT)),
            ("Primární identita", "AEGIS, Protokol: Aegis, Ultimate OS"),
            ("Charakter dokumentu", "T=0 knowledge: skutečný stav implementace, ne budoucí slib."),
        ],
        widths=(4.0, 12.4),
    )

    callout(
        doc,
        "Jedna věta",
        "AEGIS je local-first obranná vrstva pro web, která detekuje manipulační vzorce v rozhraní, umí bezpečně zasáhnout do DOMu a dává uživateli možnost zásah vrátit.",
    )

    heading(doc, "1. Produktová definice", 1)
    p(
        doc,
        "AEGIS je prototyp nové kategorie cognitive defense software. Neřeší pouze reklamu nebo tracking, ale samotné rozhodovací prostředí webu: nátlakové texty, asymetrické cookie flow, urgency efekty a prvky navržené pro manipulaci pozornosti.",
    )
    bullet(doc, "Ultimate OS je prezentační a brandový surface projektu.")
    bullet(doc, "Protokol: Aegis je produktová větev zaměřená na AI obranu webového prostředí.")
    bullet(doc, "AEGIS Results je analytická vrstva pro modelové metriky efektivity, rizik a návratnosti.")

    heading(doc, "2. Implementované části v T=0", 1)
    table_key_values(
        doc,
        [
            (
                "Web a shell",
                "Next.js 16 + TypeScript, Once UI, vlastní design shell, route pro Ultimate OS i AEGIS sekce.",
            ),
            (
                "AEGIS UI",
                "App-like ops console: threat feed, security vault, object inspector, forensic assistant a scan progress.",
            ),
            (
                "Metriky",
                "Expected Results je vnořené do AEGIS experience přes /ai-guard#expected-results a má také fallback route.",
            ),
            (
                "SEO a brand",
                "Metadata, favicon pack, manifest, robots, sitemap a tři OG obrázky pro Ultimate, Protokol: Aegis a Results.",
            ),
            (
                "Dokumentace",
                "Investor brief, technical brief, Q&A, prezentační script, slovníky pojmů a PDF podklady.",
            ),
            (
                "Vercel",
                "Projekt je nasazený na doméně aegis.d-international.eu s Next.js framework konfigurací.",
            ),
        ],
    )

    heading(doc, "3. Browser extension vrstva", 1)
    table_key_values(
        doc,
        [
            ("Manifest", "Manifest V3 extension scaffold v E:\\guard\\extension."),
            ("Content script", "Skenuje textové prvky, detekuje urgency, scarcity a social proof patterny."),
            ("Background", "Drží session data per-tab, inference výsledky a messaging mezi vrstvami."),
            ("Sidepanel", "Zobrazuje threats, inference update, audit trail, manual scan a Revert DOM."),
            ("Ikony", "Extension používá finální AEGIS shield sadu 16 / 48 / 128."),
            ("Typecheck", "Extension vrstva má vlastní tsconfig.extension.json a běží pod @ts-check."),
        ],
    )

    heading(doc, "4. Intervenční schopnosti", 1)
    bullet(doc, "Rewrite: tlakové texty se přepisují do neutrálnější podoby a originál se drží v data-* metadatech.")
    bullet(doc, "Suppression: urgency prvky dostávají jemné ztlumení animace, glow a agresivního stylingu.")
    bullet(doc, "Cookie override: consent copy se normalizuje do srozumitelnější podoby bez automatického odkliknutí.")
    bullet(doc, "Revert DOM: sidepanel umí vrátit zásahy zpět a ponechat element do reloadu bez další intervence.")

    heading(doc, "5. Lokální inference", 1)
    p(
        doc,
        "Implementovaný endpoint /api/aegis/evaluate slouží jako lokální vyhodnocovací vrstva. Bridge má default na localhost a podporuje fail-soft režim: pokud lokální model není dostupný, systém vrací local policy fallback místo pádu.",
    )
    table_key_values(
        doc,
        [
            ("Endpoint", "E:\\guard\\app\\api\\aegis\\evaluate\\route.ts"),
            ("Bridge", "E:\\guard\\extension\\inference-bridge.js"),
            ("Kontrakty", "E:\\guard\\src\\lib\\ai-guard\\contracts.ts"),
            ("Provider logika", "Ollama / custom / disabled, podle konfigurace."),
        ],
    )

    heading(doc, "6. Runtime flow", 1)
    numbered(doc, "Content script přečte stránku a vytvoří ThreatRecord[] z podezřelých prvků.")
    numbered(doc, "Background worker uloží stav pro aktivní tab a informuje sidepanel.")
    numbered(doc, "Inference bridge volitelně pošle data na lokální endpoint.")
    numbered(doc, "AEGIS aplikuje zásah: flag, rewrite, suppression nebo cookie normalization.")
    numbered(doc, "Sidepanel ukáže výsledek a umožní Revert DOM.")

    heading(doc, "7. Klíčové cesty v repozitáři", 1)
    table_key_values(
        doc,
        [
            ("Hlavní shell", "E:\\guard\\src\\App.tsx"),
            ("AEGIS sekce", "E:\\guard\\src\\components\\architecture\\SectionPages.tsx"),
            ("Ops console", "E:\\guard\\src\\components\\architecture\\AiGuardOpsConsole.tsx"),
            ("Metriky", "E:\\guard\\src\\components\\analytics\\ExpectedResultsPage.tsx"),
            ("Menu data", "E:\\guard\\src\\mocks\\megaMenuData.ts"),
            ("Extension", "E:\\guard\\extension\\content.js, background.js, sidepanel.js"),
            ("SEO", "E:\\guard\\app\\resources\\seo.ts, layout.tsx, sitemap.ts, robots.ts"),
            ("OG builder", "E:\\guard\\scripts\\build_og_images.py"),
        ],
        widths=(4.4, 12.0),
    )

    heading(doc, "8. Demo mapa", 1)
    table_key_values(
        doc,
        [
            ("Ultimate OS", "https://aegis.d-international.eu/"),
            ("Protokol: Aegis", "https://aegis.d-international.eu/ai-guard"),
            ("Philosophy", "/ai-guard/philosophy"),
            ("Detection", "/ai-guard/detection"),
            ("Defense", "/ai-guard/defense"),
            ("Build", "/ai-guard/build"),
            ("Results", "/ai-guard#expected-results"),
        ],
        widths=(4.4, 12.0),
    )

    heading(doc, "9. Co v T=0 netvrdit jako hotové", 1)
    bullet(doc, "AEGIS není ještě enterprise-ready browser security produkt.")
    bullet(doc, "Product-market fit zatím není validovaný na trhu.")
    bullet(doc, "Precision/recall modelu a dataset manipulačních patternů ještě nejsou uzavřené.")
    bullet(doc, "Unpacked Chrome flow má být ještě formálně ověřený end-to-end.")
    bullet(doc, "Část UI je app-like simulace a čeká na reálný runtime threat feed.")

    heading(doc, "10. Příkazy pro vývoj", 1)
    table_key_values(
        doc,
        [
            ("Dev server", "npm run dev"),
            ("Build", "npm run build"),
            ("Typecheck web", "npm run typecheck:web"),
            ("Typecheck extension", "npm run typecheck:extension"),
            ("Celý typecheck", "npm run typecheck"),
        ],
        widths=(4.4, 12.0),
    )

    heading(doc, "11. T=0 znalostní věty", 1)
    bullet(doc, "AEGIS už má produkční web, brand, SEO a nasazenou doménu.")
    bullet(doc, "Extension vrstva už má první reálné DOM zásahy a revert flow.")
    bullet(doc, "Local inference endpoint existuje a má fallback, takže demo nestojí na dostupnosti modelu.")
    bullet(doc, "Architektura je typovaná a připravená na rozšíření intervention modes.")
    bullet(doc, "Další nejlepší krok je end-to-end ověření extension flow a napojení reálného threat feedu.")

    out = OUTDIR / "AEGIS_knowledge_T0.docx"
    OUTDIR.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def render_docx(docx_path: Path, outdir: Path) -> None:
    if outdir.exists():
        shutil.rmtree(outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [
            str(PYTHON),
            str(RENDER_SCRIPT),
            str(docx_path),
            "--output_dir",
            str(outdir),
        ],
        text=True,
        capture_output=True,
    )
    if completed.returncode != 0:
        raise subprocess.CalledProcessError(
            completed.returncode,
            completed.args,
            output=completed.stdout,
            stderr=completed.stderr,
        )


def images_to_pdf(image_dir: Path, pdf_path: Path) -> None:
    images = sorted(image_dir.glob("page-*.png"))
    if not images:
        raise RuntimeError(f"No rendered pages in {image_dir}")

    c = None
    for image_path in images:
        img = Image.open(image_path)
        width, height = img.size
        if c is None:
            c = canvas.Canvas(str(pdf_path), pagesize=(width, height))
        c.setPageSize((width, height))
        c.drawImage(ImageReader(img), 0, 0, width=width, height=height)
        c.showPage()
        img.close()
    if c is None:
        raise RuntimeError("PDF canvas was not initialized")
    c.save()


def main() -> None:
    docx_path = build_docx()
    render_dir = OUTDIR / "render-knowledge-t0"
    pdf_path = OUTDIR / "AEGIS_knowledge_T0.pdf"
    try:
        render_docx(docx_path, render_dir)
        images_to_pdf(render_dir, pdf_path)
    except (FileNotFoundError, subprocess.CalledProcessError) as exc:
        from build_knowledge_t0_direct_pdf import main as build_direct_pdf

        print(f"DOCX renderer unavailable, falling back to direct PDF builder: {exc}")
        build_direct_pdf()
    print(docx_path)
    print(pdf_path)


if __name__ == "__main__":
    main()
