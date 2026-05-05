from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


ROOT = Path(r"E:\guard")
OUTDIR = ROOT / "dist" / "presentation-docs"
RENDER_SCRIPT = Path(
    r"C:\Users\david\.codex\plugins\cache\openai-primary-runtime\documents\26.423.10653\skills\documents\render_docx.py"
)
PYTHON = Path(
    r"C:\Users\david\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe"
)

NAVY = RGBColor(28, 39, 64)
TEAL = RGBColor(16, 128, 140)
MUTED = RGBColor(98, 108, 128)
BODY = RGBColor(32, 37, 48)


ABBREVIATIONS = [
    ("AI", "Artificial Intelligence, umělá inteligence."),
    ("LLM", "Large Language Model, velký jazykový model."),
    ("MVP", "Minimum Viable Product, nejmenší funkční verze produktu."),
    ("UX", "User Experience, uživatelský zážitek."),
    ("UI", "User Interface, uživatelské rozhraní."),
    ("API", "Rozhraní pro komunikaci mezi systémy."),
    ("DOM", "Vnitřní struktura webové stránky."),
    ("SEO", "Optimalizace pro vyhledávače."),
    ("ROI", "Návratnost investice."),
    ("KPI", "Klíčová metrika výkonu."),
    ("GDPR", "Pravidla ochrany osobních údajů v EU."),
    ("MV3", "Manifest V3, architektura browser extensions v Chrome."),
]

FOREIGN_TERMS = [
    ("local-first", "Řešení běží primárně lokálně, ne v cizím cloudu."),
    ("on-prem", "Provoz na vlastním zařízení nebo infrastruktuře."),
    ("cognitive defense", "Obrana proti digitální manipulaci a nátlaku."),
    ("dark patterns", "Manipulační designové techniky v rozhraní."),
    ("tracking", "Sledování uživatele a jeho chování."),
    ("rewrite", "Přepsání nátlakového textu do neutrální podoby."),
    ("suppression", "Ztlumení nebo omezení agresivního prvku."),
    ("revert", "Vrácení zásahu zpět."),
    ("audit trail", "Dohledatelný záznam, co systém udělal."),
    ("policy engine", "Pravidlová vrstva rozhodování bez AI modelu."),
    ("fallback", "Bezpečná náhradní cesta."),
    ("fail-soft", "Systém nespadne, jen přejde do omezeného režimu."),
]

IT_SLANG = [
    ("stack", "Soubor technologií, na kterých produkt běží."),
    ("endpoint", "Konkrétní adresa API služby."),
    ("payload", "Data posílaná mezi systémy."),
    ("scaffold", "Základní kostra projektu."),
    ("runtime", "Prostředí, ve kterém kód běží."),
    ("sidepanel", "Postranní panel rozšíření."),
    ("feed", "Průběžný tok událostí nebo záznamů."),
    ("typed contracts", "Přesně definovaná komunikace mezi částmi systému."),
    ("deploy", "Nasazení aplikace do produkce."),
    ("build", "Sestavení aplikace do spustitelné verze."),
    ("branch", "Samostatná větev repozitáře."),
    ("hotfix", "Rychlá oprava problému."),
]

HUMAN_TRANSLATIONS = [
    ("local-first AI", "AI běží hlavně u uživatele."),
    ("dark patterns", "Manipulační prvky v rozhraní."),
    ("rewrite mode", "Bezpečné přepsání nátlakového textu."),
    ("audit trail", "Záznam všech zásahů."),
    ("fallback", "Bezpečný záložní režim."),
    ("product-market fit", "Důkaz, že produkt má reálný trh."),
]


def set_font(run, size: int, color: RGBColor = BODY, bold: bool = False) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_paragraph(doc: Document, text: str, *, size: int = 11, color: RGBColor = BODY, bold: bool = False,
                  align=WD_ALIGN_PARAGRAPH.LEFT, space_after: float = 4, line_spacing: float = 1.15):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    run = p.add_run(text)
    set_font(run, size=size, color=color, bold=bold)
    return p


def add_title(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    set_font(run, 28, NAVY, True)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_after = Pt(12)
    run2 = p2.add_run(subtitle)
    set_font(run2, 11, MUTED, False)


def add_section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text.upper())
    set_font(run, 14, TEAL, True)


def add_term_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for term, explanation in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(4.2)
        cells[1].width = Cm(11.8)

        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        p0.paragraph_format.line_spacing = 1.05
        r0 = p0.add_run(term)
        set_font(r0, 10, NAVY, True)

        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        p1.paragraph_format.line_spacing = 1.08
        r1 = p1.add_run(explanation)
        set_font(r1, 10, BODY, False)

    doc.add_paragraph()


def build_docx() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    add_title(
        doc,
        "AEGIS | Slovník pojmů",
        "Tahák pro prezentaci: zkratky, cizí slova, produktové pojmy a IT slang.",
    )
    add_section_heading(doc, "Zkratky")
    add_term_table(doc, ABBREVIATIONS)

    doc.add_page_break()
    add_section_heading(doc, "Cizí slova a produktové pojmy")
    add_term_table(doc, FOREIGN_TERMS)

    doc.add_page_break()
    add_section_heading(doc, "IT slang")
    add_term_table(doc, IT_SLANG)
    add_section_heading(doc, "Jak to říct lidsky")
    add_term_table(doc, HUMAN_TRANSLATIONS)

    add_paragraph(
        doc,
        "AEGIS | aegis.d-international.eu | David Kozák | 29. 4. 2026",
        size=9,
        color=MUTED,
        space_after=0,
    )

    out = OUTDIR / "AEGIS_slovnik_pojmu_tahak.docx"
    doc.save(out)
    return out


def render_docx(docx_path: Path, outdir: Path) -> None:
    if outdir.exists():
        for child in outdir.iterdir():
            if child.is_file():
                child.unlink()
            else:
                import shutil
                shutil.rmtree(child)
    else:
        outdir.mkdir(parents=True, exist_ok=True)

    cmd = [
        str(PYTHON),
        str(RENDER_SCRIPT),
        str(docx_path),
        "--output_dir",
        str(outdir),
        "--renderer",
        "artifact-tool",
    ]
    subprocess.run(cmd, check=True)


def images_to_pdf(image_dir: Path, pdf_path: Path) -> None:
    images = sorted(image_dir.glob("page-*.png"))
    if not images:
        raise RuntimeError(f"No rendered pages in {image_dir}")

    first = Image.open(images[0])
    width, height = first.size
    first.close()

    c = canvas.Canvas(str(pdf_path), pagesize=(width, height))
    for image_path in images:
        img = Image.open(image_path)
        img_width, img_height = img.size
        c.setPageSize((img_width, img_height))
        c.drawImage(ImageReader(img), 0, 0, width=img_width, height=img_height)
        c.showPage()
        img.close()
    c.save()


def main() -> None:
    OUTDIR.mkdir(parents=True, exist_ok=True)
    docx_path = build_docx()
    render_dir = OUTDIR / "render-glossary-cheatsheet"
    pdf_path = OUTDIR / "AEGIS_slovnik_pojmu_tahak.pdf"
    render_docx(docx_path, render_dir)
    images_to_pdf(render_dir, pdf_path)
    print(docx_path)
    print(pdf_path)


if __name__ == "__main__":
    main()
