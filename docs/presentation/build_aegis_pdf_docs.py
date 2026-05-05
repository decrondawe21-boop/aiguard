from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


ROOT = Path(r"E:\guard")
TEMPLATE = Path(r"C:\Users\david\Desktop\sablona.docx")
OUTDIR = ROOT / "dist" / "presentation-docs"
RENDER_SCRIPT = Path(
    r"C:\Users\david\.codex\plugins\cache\openai-primary-runtime\documents\26.423.10653\skills\documents\render_docx.py"
)
PYTHON = Path(
    r"C:\Users\david\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe"
)

TITLE_COLOR = RGBColor(57, 67, 89)
ACCENT_COLOR = RGBColor(146, 29, 0)
BODY_COLOR = RGBColor(35, 35, 35)
MUTED_COLOR = RGBColor(84, 84, 84)


def set_run_font(run, size: int, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def clear_paragraph(paragraph) -> None:
    paragraph.clear()


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        clear_paragraph(paragraph)


def write_cell(
    cell,
    text: str,
    *,
    size: int = 12,
    bold: bool = False,
    color: RGBColor = BODY_COLOR,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    line_spacing: float = 1.15,
) -> None:
    clear_cell(cell)
    parts = text.split("\n")
    first = cell.paragraphs[0]
    first.alignment = align
    first.paragraph_format.line_spacing = line_spacing
    first.paragraph_format.space_after = Pt(4)
    run = first.add_run(parts[0] if parts else "")
    set_run_font(run, size, bold=bold, color=color)
    for part in parts[1:]:
        p = cell.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(part)
        set_run_font(run, size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def style_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(4)


def fill_title(doc: Document, title: str) -> None:
    para = doc.paragraphs[1]
    para.text = title
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in para.runs:
        set_run_font(run, 38, bold=False, color=TITLE_COLOR)


def fill_headers(table, heading: str, label_a: str, body_a: str, label_b: str, body_b: str) -> None:
    write_cell(table.cell(0, 2), heading, size=22, bold=False, color=ACCENT_COLOR)
    write_cell(table.cell(1, 0), label_a.upper(), size=10, bold=True, color=BODY_COLOR)
    write_cell(table.cell(1, 2), body_a, size=11, color=BODY_COLOR)
    write_cell(table.cell(2, 0), label_b.upper(), size=10, bold=True, color=BODY_COLOR)
    write_cell(table.cell(2, 2), body_b, size=11, color=BODY_COLOR)


def build_main_doc() -> Path:
    doc = Document(TEMPLATE)
    fill_title(doc, "AEGIS\nPodklady pro prezentaci")

    tables = doc.tables
    fill_headers(
        tables[0],
        "PROJEKT",
        "Prezentace pro",
        "Investoři a IT/developer publikum",
        "Připravil",
        "David Kozák\nAEGIS | https://aegis.d-international.eu",
    )

    fill_headers(
        tables[1],
        "SHRNUTÍ",
        "Jednověté shrnutí",
        "AEGIS je local-first AI obranná vrstva pro web, která chrání uživatele před digitální manipulací, nátlakovým rozhraním a agresivním trackingem.",
        "Problém a řešení",
        "Web je dnes optimalizovaný pro konverzi, ne pro uživatele.\nAEGIS čte stránku jako útokový povrch, detekuje manipulační patterny a umí tlakové prvky přepsat nebo ztlumit bez rozbití layoutu.",
    )

    write_cell(tables[2].cell(0, 2), "PRODUKT A MVP", size=22, color=ACCENT_COLOR)
    write_cell(tables[2].cell(1, 0), "CO JE HOTOVÉ", size=10, bold=True)
    write_cell(
        tables[2].cell(1, 2),
        "Veřejný web a produktové sekce AEGIS\nBrowser extension scaffold v Manifest V3\nPrvní reálné DOM zásahy: rewrite, suppression, cookie override\nLokální inference endpoint a Revert DOM",
        size=10,
    )
    write_cell(tables[2].cell(2, 0), "DALŠÍ KROK", size=10, bold=True)
    write_cell(
        tables[2].cell(2, 2),
        "Ověřit unpacked Chrome flow end-to-end\nDopojit real threat feed a audit trail\nVylepšit inference kvalitu\nPřipravit distribuovatelnou browser verzi",
        size=10,
    )

    write_cell(
        tables[3].cell(0, 0),
        "INVESTIČNÍ\nTEZE",
        size=10,
        bold=True,
    )
    write_cell(
        tables[3].cell(0, 2),
        "AEGIS není jen další adblock. Vytváří novou kategorii: cognitive defense software. Klasické nástroje blokují bannery nebo trackery, ale AEGIS řeší samotný úmysl rozhraní a ochranu úsudku uživatele.",
        size=11,
    )
    write_cell(tables[3].cell(1, 0), "OČEKÁVANÝ\nDOPAD", size=10, bold=True)
    write_cell(
        tables[3].cell(1, 2),
        "Modelová projekce: +145 % efektivita | -88 % kritická rizika | 4,5 měsíce break-even",
        size=10,
    )
    write_cell(tables[3].cell(2, 0), "TECHNICKÁ\nVRSTVA", size=10, bold=True)
    write_cell(
        tables[3].cell(2, 2),
        "Next.js 16 + TypeScript + browser extension + local endpoint",
        size=10,
    )
    write_cell(tables[3].cell(3, 0), "RIZIKA A\nDALŠÍ KROKY", size=10, bold=True)
    write_cell(
        tables[3].cell(3, 2),
        "Local-first inference, revert flow a audit trail\nPotvrdit product-market fit a přesnost zásahů\nPilotní validace a governance vrstva",
        size=10,
    )
    write_cell(
        tables[3].cell(4, 2),
        "AEGIS | aegis.d-international.eu | David Kozák | Podklady pro prezentaci | 29. 4. 2026",
        size=10,
        color=MUTED_COLOR,
    )

    for table in tables:
        style_table(table)

    out = OUTDIR / "AEGIS_podklady_prezentace.docx"
    doc.save(out)
    return out


def build_glossary_doc() -> Path:
    doc = Document(TEMPLATE)
    fill_title(doc, "AEGIS\nSlovník cizích slov")

    tables = doc.tables
    fill_headers(
        tables[0],
        "KONTEXT",
        "Použití",
        "Krátký tahák pro prezentaci investorům a IT publiku.",
        "Pravidlo",
        "Když mluvíš naživo, vysvětluj termíny lidsky.",
    )

    fill_headers(
        tables[1],
        "ZÁKLADNÍ POJMY",
        "Cognitive defense",
        "Obrana proti digitální manipulaci a nátlaku na rozhodování.",
        "Local-first AI",
        "AI běží primárně lokálně, takže data nemusí odcházet do cizího cloudu.",
    )

    write_cell(tables[2].cell(0, 2), "TECHNICKÉ POJMY", size=22, color=ACCENT_COLOR)
    write_cell(tables[2].cell(1, 0), "INFERENCE | DOM |\nSIDEPANEL |\nTHREAT FEED", size=10, bold=True)
    write_cell(
        tables[2].cell(1, 2),
        "Inference = vyhodnocení stránky nebo textu modelem\nDOM = vnitřní struktura webové stránky\nSidepanel = postranní panel rozšíření\nThreat feed = živý přehled zachycených hrozeb",
        size=11,
    )
    write_cell(tables[2].cell(2, 0), "REWRITE |\nSUPPRESSION |\nREVERT |\nAUDIT TRAIL", size=10, bold=True)
    write_cell(
        tables[2].cell(2, 2),
        "Rewrite = přepsání nátlakového textu do neutrálnější podoby\nSuppression = ztlumení tlakového vizuálního efektu\nRevert = vrácení zásahu zpět\nAudit trail = záznam toho, co systém detekoval a co udělal",
        size=11,
    )

    write_cell(tables[3].cell(0, 0), "PRODUKTOVÉ A\nBYZNYS POJMY", size=10, bold=True)
    write_cell(
        tables[3].cell(0, 2),
        "Policy engine = pravidlová vrstva bez AI modelu.",
        size=9,
    )
    write_cell(tables[3].cell(1, 0), "INVESTOR\nSLOVA", size=10, bold=True)
    write_cell(
        tables[3].cell(1, 2),
        "Moat = obrana před konkurencí.",
        size=9,
    )
    write_cell(tables[3].cell(2, 0), "DEVELOPER\nSLOVA", size=10, bold=True)
    write_cell(
        tables[3].cell(2, 2),
        "Fallback = bezpečná náhradní cesta.",
        size=9,
    )
    write_cell(tables[3].cell(3, 0), "SILNÉ\nVĚTY", size=10, bold=True)
    write_cell(
        tables[3].cell(3, 2),
        "AEGIS chrání úsudek uživatele.\nKaždý zásah je vratný.",
        size=9,
    )
    write_cell(
        tables[3].cell(4, 2),
        "AEGIS | Slovník cizích slov | David Kozák | 29. 4. 2026",
        size=10,
        color=MUTED_COLOR,
    )

    for table in tables:
        style_table(table)

    out = OUTDIR / "AEGIS_slovnik_cizich_slov.docx"
    doc.save(out)
    return out


def render_docx(docx_path: Path, outdir: Path) -> None:
    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [str(PYTHON), str(RENDER_SCRIPT), str(docx_path), "--output_dir", str(outdir), "--renderer", "artifact-tool"],
        check=True,
    )


def pngs_to_pdf(png_dir: Path, pdf_path: Path) -> None:
    pages = sorted(png_dir.glob("page-*.png"))
    if not pages:
        raise RuntimeError(f"No PNG pages found in {png_dir}")
    first_image = Image.open(pages[0])
    width, height = first_image.size
    first_image.close()
    pdf = canvas.Canvas(str(pdf_path), pagesize=(width, height))
    for page in pages:
        image = Image.open(page)
        w, h = image.size
        image.close()
        pdf.setPageSize((w, h))
        pdf.drawImage(ImageReader(str(page)), 0, 0, width=w, height=h)
        pdf.showPage()
    pdf.save()


def main() -> None:
    if OUTDIR.exists():
        shutil.rmtree(OUTDIR)
    OUTDIR.mkdir(parents=True, exist_ok=True)

    main_doc = build_main_doc()
    glossary_doc = build_glossary_doc()

    main_render = OUTDIR / "render-main"
    glossary_render = OUTDIR / "render-glossary"

    render_docx(main_doc, main_render)
    render_docx(glossary_doc, glossary_render)

    pngs_to_pdf(main_render, OUTDIR / "AEGIS_podklady_prezentace.pdf")
    pngs_to_pdf(glossary_render, OUTDIR / "AEGIS_slovnik_cizich_slov.pdf")

    print(main_doc)
    print(glossary_doc)


if __name__ == "__main__":
    main()
